"""Tests για το star schema: κάθε «κλειδί» των facts πρέπει να υπάρχει στη διάστασή του."""
import pytest
from openpyxl import load_workbook

from audit_ai import database as db
from audit_ai import generator, pipeline, powerbi_export


@pytest.fixture(scope="module")
def env(tmp_path_factory):
    folder = tmp_path_factory.mktemp("audit")
    gl, tb, truth = generator.generate()
    paths = generator.save(gl, tb, truth, raw_dir=folder)
    conn = db.connect(pipeline.run(paths["gl"], paths["tb"], folder / "audit.db", folder / "r.xlsx"))
    return conn, folder


@pytest.fixture(scope="module")
def tables(env):
    conn, folder = env
    return powerbi_export.export(conn, folder / "powerbi")


def test_every_line_has_a_user_account_and_date(tables):
    facts = tables["fact_journal_lines"]
    assert set(facts["created_by"]) <= set(tables["dim_user"]["created_by"])
    assert set(facts["account"]) <= set(tables["dim_account"]["account"])
    assert set(facts["posting_date"]) <= set(tables["dim_date"]["date"])


def test_dimensions_have_unique_keys(tables):
    assert tables["dim_user"]["created_by"].is_unique
    assert tables["dim_account"]["account"].is_unique
    assert tables["dim_date"]["date"].is_unique


def test_high_risk_flag_matches_the_list(tables):
    flagged = tables["fact_journal_lines"].query("is_high_risk == 1")["journal_id"].nunique()
    assert flagged == 19


def test_roles_are_assigned(tables):
    roles = tables["dim_user"].set_index("created_by")["role"]
    assert roles["CFO"] == "Διοίκηση" and roles["TEMP_USER"] == "Μη εξουσιοδοτημένος"


def test_excel_has_one_named_table_per_sheet(env):
    conn, folder = env
    wb = load_workbook(powerbi_export.export_excel(conn, folder / "model.xlsx"))
    for sheet in wb.worksheets:
        assert list(sheet.tables) == [sheet.title]
