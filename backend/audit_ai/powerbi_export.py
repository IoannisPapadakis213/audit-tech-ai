"""Εξαγωγή για Power BI: star schema σε CSV και σε ένα Excel (για το Power BI στον browser).

Star schema = στο κέντρο οι πίνακες γεγονότων (facts: «τι συνέβη»), γύρω οι διαστάσεις
(dimensions: «ποιος, πότε, πού»). Το Power BI δουλεύει καλύτερα, πιο γρήγορα και πιο απλά έτσι.

    dim_date ──┐                ┌── dim_user
               ├── fact_journal_lines ──┤
  dim_account ─┘        │       └── (benford ανά χρήστη μέσα στο dim_user)
                  fact_rule_hits
    dim_digit (Benford)    trial_balance (συμφωνία GL-TB)
"""
import pandas as pd
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

from audit_ai import config as cfg
from audit_ai import database as db


def user_role(user):
    if user in cfg.SENIOR_MANAGEMENT:
        return "Διοίκηση"
    if user in cfg.SYSTEM_USERS:
        return "Σύστημα"
    if user in cfg.AUTHORIZED_POSTERS:
        return "Εξουσιοδοτημένος"
    return "Μη εξουσιοδοτημένος"


def build_star_schema(conn):
    """Επιστρέφει dict {όνομα πίνακα: DataFrame}."""
    lines = db.query(conn, """
        SELECT g.journal_id, g.line_id, DATE(g.posting_date) AS posting_date, g.entry_datetime,
               CAST(strftime('%H', g.entry_datetime) AS INTEGER) AS entry_hour,
               CAST(strftime('%w', g.entry_datetime) AS INTEGER) AS entry_weekday_sun0,
               g.account, g.debit, g.credit, g.amount, g.created_by, g.approved_by, g.source,
               g.description, g.vendor_id, g.is_weekend, g.is_holiday, g.is_after_hours,
               g.is_round_amount, g.is_manual, g.days_to_year_end,
               COALESCE(s.risk_score, 0) AS risk_score,
               CASE WHEN s.risk_score >= ? THEN 1 ELSE 0 END AS is_high_risk
        FROM silver_gl AS g LEFT JOIN risk_scores AS s ON s.journal_id = g.journal_id
    """, (cfg.RISK_SCORE_THRESHOLD,))
    # Δευτέρα = 1 ... Κυριακή = 7, για τον χάρτη θερμότητας ημέρα × ώρα
    lines["entry_weekday"] = lines.pop("entry_weekday_sun0").replace(0, 7)

    rule_hits = db.query(conn, "SELECT journal_id, rule, weight FROM rule_hits")

    dates = pd.DataFrame({"date": pd.date_range(cfg.FISCAL_YEAR_START, cfg.POST_CLOSE_END)})
    holidays = pd.to_datetime(list(cfg.GREEK_HOLIDAYS_2025))
    dim_date = dates.assign(
        year=dates["date"].dt.year, month=dates["date"].dt.month,
        month_name=dates["date"].dt.strftime("%Y-%m"), weekday=dates["date"].dt.dayofweek + 1,
        is_weekend=dates["date"].dt.dayofweek >= 5, is_holiday=dates["date"].isin(holidays),
        is_post_close=dates["date"] > pd.Timestamp(cfg.FISCAL_YEAR_END),
    )
    dim_date["date"] = dim_date["date"].dt.strftime("%Y-%m-%d")

    users = sorted(set(lines["created_by"]))
    benford = db.query(conn, "SELECT created_by, n AS benford_n, mad AS benford_mad, conformity AS benford_conformity FROM benford_by_user")
    dim_user = (pd.DataFrame({"created_by": users})
                .assign(role=lambda d: d["created_by"].map(user_role))
                .merge(benford, on="created_by", how="left"))

    uses = lines.groupby("account")["journal_id"].nunique()
    dim_account = pd.DataFrame(
        [{"account": acc, "account_name": name, "nature": nature, "journals": int(uses.get(acc, 0))}
         for acc, (name, nature) in cfg.CHART_OF_ACCOUNTS.items()])

    dim_digit = db.query(conn, "SELECT digit, expected, actual, count FROM benford_overall")
    trial_balance = db.query(conn, "SELECT * FROM dq_reconciliation")

    return {"fact_journal_lines": lines, "fact_rule_hits": rule_hits, "dim_date": dim_date,
            "dim_user": dim_user, "dim_account": dim_account, "dim_digit": dim_digit,
            "trial_balance": trial_balance}


def export(conn, out_dir=cfg.POWERBI_DIR):
    out_dir.mkdir(parents=True, exist_ok=True)
    tables = build_star_schema(conn)
    for name, df in tables.items():
        df.to_csv(out_dir / f"{name}.csv", index=False, encoding="utf-8-sig")   # utf-8-sig: σωστά ελληνικά στο Power BI/Excel
    return tables


def export_excel(conn, path=None):
    """Ένα Excel, ένα φύλλο ανά πίνακα, και κάθε φύλλο ως «Excel Table».
    Το Power BI στον browser διαβάζει κάθε Excel Table ως ξεχωριστό πίνακα του μοντέλου."""
    path = path or cfg.POWERBI_DIR / "audit_model.xlsx"
    path.parent.mkdir(parents=True, exist_ok=True)
    tables = build_star_schema(conn)
    with pd.ExcelWriter(path, engine="openpyxl") as xl:
        for name, df in tables.items():
            df.to_excel(xl, sheet_name=name, index=False)
            sheet = xl.sheets[name]
            ref = f"A1:{get_column_letter(df.shape[1])}{len(df) + 1}"
            table = Table(displayName=name, ref=ref)
            table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
            sheet.add_table(table)
    return path


if __name__ == "__main__":
    with db.connect() as conn:
        tables = export(conn)
        xlsx = export_excel(conn)
    for name, df in tables.items():
        print(f"{name:20s} {len(df):6d} γραμμές")
    print(f"\nΦάκελος -> {cfg.POWERBI_DIR}")
    print(f"Excel για Power BI -> {xlsx}")
